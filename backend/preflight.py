"""导出前体检 —— 在拼页阶段就把「导出后才会发现的错」提前摆出来。

检查项全部是确定性的（不猜）：
  1. 源文件不见了 —— 有缩略图的页会降级成整页图，没缩略图的页直接丢
  2. 源文件被改过（磁盘大小/修改时间与库里记录不一致）—— 缩略图和页码可能已经过期，建议回「扫描处理」重扫
  3. 页面比例混排（4:3 与 16:9 拼在一份稿里）—— 导出的文件会一页一个尺寸
  4. 会被降级成整页图的页（含图表/SmartArt/嵌入对象，或旧版 .ppt）—— 图不能改字
  5. 同一页在稿件里出现多次

字体缺失这类"猜"的检查故意没做：字体名和字体文件名不是一回事，误报比漏报更烦人。
"""
import os

from pptx import Presentation


def _ratio_label(w, h):
    """把 slide_width/height 折成 16:9 / 4:3 / 其他。"""
    if not w or not h:
        return '未知'
    r = w / h
    known = [(16 / 9, '16:9'), (4 / 3, '4:3'), (16 / 10, '16:10')]
    for val, label in known:
        if abs(r - val) < 0.01:
            return label
    return f'{r:.2f}:1'


def check(db, items, risky_kinds=None):
    """items: get_composition_items() 的行。返回结构化体检报告。"""
    if risky_kinds is None:
        from composer import _risky_kinds as risky_kinds

    issues = []
    per_page = []
    size_cache = {}       # file_path -> {'ratio':..,'width':..,'height':..,'prs':Presentation|None}
    missing_pages = []    # 源文件不在
    changed_pages = []
    degraded_pages = []
    dup_counter = {}

    for idx, it in enumerate(items, start=1):
        path = it['file_path'] or ''
        num = it['slide_number']
        label = {
            'index': idx,
            'title': it['title'] or '',
            'file_name': it['file_name'] or '',
            'slide_number': num,
            'status': 'ok',
            'reason': '',
        }
        dup_key = (path, num)
        dup_counter[dup_key] = dup_counter.get(dup_key, 0) + 1

        if not path or not os.path.exists(path):
            label['status'] = 'missing'
            label['reason'] = '源文件不在'
            missing_pages.append(idx)
            per_page.append(label)
            continue

        # 源文件是否被改过
        rec = db.get_file_by_path(path) if hasattr(db, 'get_file_by_path') else None
        if rec:
            try:
                st = os.stat(path)
                if (rec['modified_time'] and abs(st.st_mtime - rec['modified_time']) > 1) \
                        or (rec['file_size'] and st.st_size != rec['file_size']):
                    changed_pages.append(idx)
            except OSError:
                pass

        # 页面尺寸（每个源文件只开一次）
        if path not in size_cache:
            info = {'ratio': None, 'width': None, 'height': None, 'prs': None}
            if path.lower().endswith('.pptx'):
                try:
                    prs = Presentation(path)
                    info['prs'] = prs
                    info['width'] = int(prs.slide_width or 0)
                    info['height'] = int(prs.slide_height or 0)
                    info['ratio'] = _ratio_label(info['width'], info['height'])
                except Exception:
                    info['ratio'] = '打不开'
            else:
                info['ratio'] = '.ppt 旧格式'
            size_cache[path] = info
        sinfo = size_cache[path]
        label['ratio'] = sinfo['ratio']

        # 会不会降级成整页图
        reason = ''
        prs = sinfo['prs']
        if prs is None:
            reason = '.ppt 旧格式' if path.lower().endswith('.ppt') else '源文件打不开'
        elif not (1 <= num <= len(prs.slides)):
            reason = f'源文件只剩 {len(prs.slides)} 页，第 {num} 页已不存在'
        else:
            try:
                kinds = risky_kinds(prs.slides[num - 1])
            except Exception:
                kinds = set()
            if kinds:
                reason = '、'.join(sorted(kinds))
        if reason:
            label['status'] = 'degraded'
            label['reason'] = reason
            degraded_pages.append(idx)

        per_page.append(label)

    # ── 汇总成 issues ──
    if missing_pages:
        issues.append({
            'level': 'blocker',
            'code': 'missing_file',
            'title': f'{len(missing_pages)} 页的源文件找不到了',
            'detail': '这一页只能用缩略图顶上（变成不可编辑的整页图）；没有缩略图的页会直接丢掉。'
                      '把源文件放回原路径，或重新扫描导入一次。',
            'pages': missing_pages[:60],
        })
    if changed_pages:
        issues.append({
            'level': 'warn',
            'code': 'file_changed',
            'title': f'{len(changed_pages)} 页的源文件在入库后被改过',
            'detail': '缩略图、标题和页码可能还是旧的。建议先到「扫描处理」重扫这几份文件再导出。',
            'pages': changed_pages[:60],
        })
    if degraded_pages:
        issues.append({
            'level': 'warn',
            'code': 'degraded',
            'title': f'{len(degraded_pages)} 页会变成整页图（导出后不能改字）',
            'detail': '含图表、SmartArt 或嵌入对象的页没法可靠搬运，只能整页截图；来源会写进备注。',
            'pages': degraded_pages[:60],
        })

    ratio_pages = {}
    for p in per_page:
        lbl = p.get('ratio') or '未知'
        ratio_pages[lbl] = ratio_pages.get(lbl, 0) + 1
    files_by_ratio = {}
    for path, info in size_cache.items():
        files_by_ratio.setdefault(info['ratio'], []).append(os.path.basename(path))
    ratio_rows = [{'label': k, 'pages': v, 'files': files_by_ratio.get(k, [])[:6]}
                  for k, v in ratio_pages.items()]
    real_ratios = [r for r in ratio_rows if r['label'] in ('16:9', '4:3', '16:10')]
    if len(real_ratios) > 1:
        issues.append({
            'level': 'warn',
            'code': 'mixed_ratio',
            'title': f'页面比例混排：{" / ".join(r["label"] for r in real_ratios)}',
            'detail': '一份 PPT 里页面尺寸不统一，导出后会有页比别的页宽或矮一截。'
                      '最好只用同一比例的页，或在 PowerPoint 里统一版式后再导出。',
            'pages': [],
        })

    dups = [k for k, v in dup_counter.items() if v > 1]
    if dups:
        issues.append({
            'level': 'info',
            'code': 'duplicate',
            'title': f'{len(dups)} 页在稿件里重复出现',
            'detail': '同一页被加了多次。想留就留（同一页本来就允许多次），只是提醒一下。',
            'pages': [],
        })

    return {
        'ok': not any(i['level'] == 'blocker' for i in issues),
        'pages': len(items),
        'files': len(size_cache),
        'issues': issues,
        'ratios': ratio_rows,
        'per_page': per_page,
    }
