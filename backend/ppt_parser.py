import os
import sys
import shutil
import hashlib
import logging

from logger import get_logger

_log = get_logger(__name__)


def extract_text_from_slide(slide):
    """Extract all text content from a python-pptx slide object."""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        texts.append(t)
    return '\n'.join(texts)


def extract_title_from_slide(slide):
    """Try to extract the title from a slide."""
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text.strip()
        if title:
            return title
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                return t[:100]
    return ''


def compute_content_hash(text):
    normalized = ''.join(text.lower().split())
    return hashlib.md5(normalized.encode('utf-8')).hexdigest()


class PartialParseError(Exception):
    """文件只被部分解析。携带已解析的部分页面，供调用方决定如何处理。

    用于区分"完全失败"（返回空）与"部分成功"（拿到一部分页面但不完整），
    避免把不完整的解析结果静默当作完整数据入库。
    """
    def __init__(self, slides_data, message):
        super().__init__(message)
        self.slides_data = slides_data


def count_slides_fast(file_path):
    """快速数页数，不解析内容。

    .pptx / .pptm 直接读包里的 presentation.xml，数 sldIdLst 里真正挂在根元素下的
    sldId —— 毫秒级，用来在导入前算出总页数（进度条按页走）。
    不能用「整个文件里 sldId 出现几次」：带 AlternateContent 的 deck 会写两份
    （Choice + Fallback），数出来正好翻倍。老 .ppt 二进制要转换才知道，返回 None。
    """
    import xml.etree.ElementTree as ET
    import zipfile

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in ('.pptx', '.pptm'):
        return None
    try:
        with zipfile.ZipFile(file_path) as z:
            xml = z.read('ppt/presentation.xml')
        root = ET.fromstring(xml)
        ns_p = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
        lst = root.find(ns_p + 'sldIdLst')      # 只认根元素下的那个
        if lst is None:
            return None
        n = len(list(lst.findall(ns_p + 'sldId')))
        return n if n > 0 else None
    except Exception as e:
        _log.debug('count_slides_fast failed for %s: %s', file_path, e)
        return None


def parse_pptx(file_path):
    """Parse a .pptx file using python-pptx. Returns list of slide dicts.
    Raises PartialParseError if the file was only partially parsed."""
    from pptx import Presentation
    slides_data = []
    total = 0
    try:
        prs = Presentation(file_path)
        total = len(prs.slides)
        for i, slide in enumerate(prs.slides):
            text_content = extract_text_from_slide(slide)
            title = extract_title_from_slide(slide)
            content_hash = compute_content_hash(text_content)
            slides_data.append({
                'slide_number': i + 1,
                'title': title,
                'text_content': text_content,
                'content_hash': content_hash,
            })
    except Exception as e:
        if slides_data:
            # Mid-parse failure: keep partial results but signal incompleteness
            _log.error("Partial parse of pptx %s (%d/%d slides): %s",
                       file_path, len(slides_data), total, e)
            raise PartialParseError(
                slides_data,
                f"文件未完整解析（仅 {len(slides_data)}/{total} 页）：{e}")
        _log.error("Error parsing pptx %s: %s", file_path, e)
    return slides_data


def parse_ppt_with_com(file_path):
    """Parse old .ppt format using PowerPoint COM automation."""
    from com_utils import powerpoint_app, open_presentation

    slides_data = []
    total = 0
    try:
        with powerpoint_app() as ppt_app:
            with open_presentation(ppt_app, file_path, read_only=True) as pres:
                total = pres.Slides.Count
                for i in range(1, pres.Slides.Count + 1):
                    slide = pres.Slides(i)
                    texts = []
                    for si in range(1, slide.Shapes.Count + 1):
                        shape = slide.Shapes(si)
                        try:
                            if shape.HasTextFrame:
                                t = shape.TextFrame.TextRange.Text.strip()
                                if t:
                                    texts.append(t)
                        except Exception:
                            pass
                        try:
                            if shape.HasTable:
                                tbl = shape.Table
                                for r in range(1, tbl.Rows.Count + 1):
                                    for c in range(1, tbl.Columns.Count + 1):
                                        ct = tbl.Cell(r, c).Shape.TextFrame.TextRange.Text.strip()
                                        if ct:
                                            texts.append(ct)
                        except Exception:
                            pass

                    text_content = '\n'.join(texts)
                    title = ''
                    try:
                        title = slide.Shapes.Title.TextFrame.TextRange.Text.strip()
                    except Exception:
                        if texts:
                            title = texts[0][:100]

                    content_hash = compute_content_hash(text_content)
                    slides_data.append({
                        'slide_number': i,
                        'title': title,
                        'text_content': text_content,
                        'content_hash': content_hash,
                    })
    except Exception as e:
        if slides_data:
            _log.error("Partial parse of ppt %s (%d/%d slides): %s",
                       file_path, len(slides_data), total, e)
            raise PartialParseError(
                slides_data,
                f"文件未完整解析（仅 {len(slides_data)}/{total} 页）：{e}")
        _log.error("Error parsing ppt %s: %s", file_path, e)
    return slides_data


def parse_ppt_with_libreoffice(file_path):
    """Parse old .ppt/.pps format on macOS/Linux: LibreOffice 转 pptx 后按 pptx 解析。

    COM 只存在于 Windows；macOS 上没有 PowerPoint 自动化，用 LibreOffice 无头转换替代。
    """
    import libreoffice

    converted, tmp_root = libreoffice.convert(file_path, 'pptx')
    try:
        return parse_pptx(converted)
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)


def parse_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pptx':
        return parse_pptx(file_path)
    elif ext in ('.ppt', '.pps'):
        if sys.platform == 'win32':
            return parse_ppt_with_com(file_path)
        return parse_ppt_with_libreoffice(file_path)
    return []


def delete_slide_from_file(file_path, slide_number):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(file_path)
        xml_slides = prs.slides._sldIdLst
        slides = list(xml_slides)
        xml_slides.remove(slides[slide_number - 1])
        prs.save(file_path)
    elif ext in ('.ppt', '.pps', '.ppsx'):
        if sys.platform != 'win32':
            # COM 只在 Windows 有；LibreOffice 转格式再存回会改变文件格式，不能默默做
            raise RuntimeError(
                '删除旧版 .ppt/.pps 里的页面需要 Windows 上的 PowerPoint；'
                'macOS 上请先另存为 .pptx 再删，或改在 Windows 端操作。')
        from com_utils import powerpoint_app, open_presentation

        with powerpoint_app() as ppt_app:
            with open_presentation(ppt_app, file_path, read_only=False) as pres:
                pres.Slides(slide_number).Delete()
                pres.Save()
